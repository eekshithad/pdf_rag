from docx import Document

from rag.retriever import Retriever
from rag.reranker import Reranker
from rag.prompt import build_prompt
from rag.llm import LLM


class ReportService:

    def __init__(self):

        self.retriever = Retriever()

        self.reranker = Reranker()

        self.llm = LLM()

    def generate_report(self, topic):

        docs = self.retriever.search(topic)

        docs = self.reranker.rerank(topic, docs)

        prompt = build_prompt(topic, docs)

        answer = self.llm.generate(prompt)

        document = Document()

        document.add_heading("AI Generated Report", level=1)

        document.add_heading("Topic", level=2)

        document.add_paragraph(topic)

        document.add_heading("Summary", level=2)

        document.add_paragraph(answer)

        document.add_heading("Sources", level=2)

        for doc in docs:

            file = doc.metadata["file_name"]

            page = doc.metadata["page"] + 1

            document.add_paragraph(
                f"{file} (Page {page})"
            )

        output = f"reports/{topic}.docx"

        document.save(output)

        return output